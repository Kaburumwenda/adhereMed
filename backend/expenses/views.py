from datetime import datetime, timedelta, date
from decimal import Decimal, InvalidOperation
import csv
import io
import uuid

from django.db.models import Count, Sum, Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, filters, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend

from .models import Expense, ExpenseCategory
from .serializers import ExpenseSerializer, ExpenseCategorySerializer


class ExpenseCategoryViewSet(viewsets.ModelViewSet):
    queryset = ExpenseCategory.objects.all()
    serializer_class = ExpenseCategorySerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['is_active']
    search_fields = ['name']
    ordering_fields = ['name', 'created_at']

    def get_queryset(self):
        qs = super().get_queryset()
        return qs.annotate(
            expense_count=Count('expenses'),
            total_spent=Sum('expenses__amount'),
        )


class ExpenseViewSet(viewsets.ModelViewSet):
    queryset = Expense.objects.select_related('category', 'supplier', 'submitted_by', 'approved_by').all()
    serializer_class = ExpenseSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['status', 'category', 'payment_method', 'is_recurring', 'supplier']
    search_fields = ['reference', 'title', 'vendor', 'description', 'payment_reference']
    ordering_fields = ['expense_date', 'amount', 'created_at', 'due_date']

    def get_queryset(self):
        qs = super().get_queryset()
        params = self.request.query_params
        date_from = params.get('date_from')
        date_to = params.get('date_to')
        min_amount = params.get('min_amount')
        max_amount = params.get('max_amount')
        if date_from:
            qs = qs.filter(expense_date__gte=date_from)
        if date_to:
            qs = qs.filter(expense_date__lte=date_to)
        if min_amount:
            try:
                qs = qs.filter(amount__gte=Decimal(min_amount))
            except Exception:
                pass
        if max_amount:
            try:
                qs = qs.filter(amount__lte=Decimal(max_amount))
            except Exception:
                pass
        return qs

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        exp = self.get_object()
        if exp.status not in (Expense.Status.PENDING,):
            return Response({'detail': f'Cannot approve from status {exp.status}.'}, status=400)
        exp.status = Expense.Status.APPROVED
        exp.approved_at = timezone.now()
        if request.user and request.user.is_authenticated:
            exp.approved_by = request.user
        exp.save(update_fields=['status', 'approved_at', 'approved_by', 'updated_at'])
        return Response(self.get_serializer(exp).data)

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        exp = self.get_object()
        reason = request.data.get('reason') or ''
        exp.status = Expense.Status.REJECTED
        if reason:
            exp.notes = (exp.notes + '\n' if exp.notes else '') + f'Rejected: {reason}'
        exp.save(update_fields=['status', 'notes', 'updated_at'])
        return Response(self.get_serializer(exp).data)

    @action(detail=True, methods=['post'])
    def mark_paid(self, request, pk=None):
        exp = self.get_object()
        if exp.status not in (Expense.Status.APPROVED, Expense.Status.PENDING):
            return Response({'detail': f'Cannot mark as paid from status {exp.status}.'}, status=400)
        exp.status = Expense.Status.PAID
        exp.paid_at = timezone.now()
        # Optional payment-method override on payment
        method = request.data.get('payment_method')
        ref = request.data.get('payment_reference')
        update_fields = ['status', 'paid_at', 'updated_at']
        if method:
            exp.payment_method = method
            update_fields.append('payment_method')
        if ref:
            exp.payment_reference = ref
            update_fields.append('payment_reference')
        exp.save(update_fields=update_fields)
        return Response(self.get_serializer(exp).data)

    @action(detail=False, methods=['get'], url_path='report-pdf')
    def report_pdf(self, request):
        """Generate a branded premium PDF report of expenses.

        Respects the same filters as the list endpoint (status, category,
        payment_method, supplier) plus date_from / date_to.
        """
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.units import mm
            from reportlab.lib.colors import HexColor
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.enums import TA_RIGHT, TA_CENTER
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
            )
            from django.utils import timezone as _tz
        except ImportError:
            return Response({'detail': 'reportlab not installed.'}, status=500)

        qs = self.filter_queryset(self.get_queryset())
        expenses = list(qs)

        from config.pdf_branding import resolve_branding
        logo_path, business_name = resolve_branding(request)

        NAVY = HexColor('#0F172A')
        TEAL = HexColor('#0D9488')
        SLATE = HexColor('#64748B')
        BORDER = HexColor('#E2E8F0')
        BG = HexColor('#F8FAFC')
        WHITE = HexColor('#FFFFFF')
        STATUS_COLORS = {
            'pending': HexColor('#D97706'),
            'approved': HexColor('#2563EB'),
            'paid': HexColor('#16A34A'),
            'rejected': HexColor('#DC2626'),
            'cancelled': HexColor('#64748B'),
        }
        STATUS_LABELS = dict(Expense.Status.choices)

        def money(v):
            try:
                return f'KSh {float(v or 0):,.2f}'
            except (TypeError, ValueError):
                return 'KSh 0.00'

        def fmt_date(d):
            return d.strftime('%d %b %Y') if d else '—'

        # ── Aggregates ────────────────────────────────────────────────────
        total_amount = 0.0
        total_tax = 0.0
        status_agg = {}
        category_agg = {}
        for e in expenses:
            amount = float(e.amount or 0)
            total_amount += amount
            total_tax += float(e.tax_amount or 0)
            key = e.status
            if key not in status_agg:
                status_agg[key] = {'count': 0, 'value': 0.0}
            status_agg[key]['count'] += 1
            status_agg[key]['value'] += amount
            cname = e.category.name if e.category else 'Uncategorised'
            category_agg[cname] = category_agg.get(cname, 0.0) + amount
        month_start = _tz.localtime().date().replace(day=1)
        this_month = sum(float(e.amount or 0) for e in expenses
                         if e.expense_date and e.expense_date >= month_start)
        pending_value = status_agg.get('pending', {}).get('value', 0)
        paid_value = status_agg.get('paid', {}).get('value', 0)
        top_categories = sorted(category_agg.items(), key=lambda x: -x[1])[:5]

        PAGE = landscape(A4)
        PAGE_W = PAGE[0] - 30 * mm

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=PAGE,
            leftMargin=15 * mm, rightMargin=15 * mm,
            topMargin=12 * mm, bottomMargin=18 * mm,
            title='Operating Expenses Report',
        )

        st_brand = ParagraphStyle('brand', fontName='Helvetica-Bold', fontSize=13,
                                   textColor=NAVY, leading=16)
        st_brand_sub = ParagraphStyle('brandsub', fontName='Helvetica', fontSize=8,
                                      textColor=SLATE, leading=11)
        st_doctitle = ParagraphStyle('doctitle', fontName='Helvetica-Bold', fontSize=20,
                                     textColor=TEAL, leading=24, alignment=TA_RIGHT)
        st_docsub = ParagraphStyle('docsub', fontName='Helvetica', fontSize=8.5,
                                   textColor=SLATE, leading=12, alignment=TA_RIGHT)
        st_label = ParagraphStyle('label', fontName='Helvetica-Bold', fontSize=6.5,
                                  textColor=SLATE, leading=9)
        st_value = ParagraphStyle('value', fontName='Helvetica-Bold', fontSize=11,
                                  textColor=NAVY, leading=14)
        st_value_sub = ParagraphStyle('valuesub', fontName='Helvetica', fontSize=7,
                                      textColor=SLATE, leading=9)
        st_section = ParagraphStyle('section', fontName='Helvetica-Bold', fontSize=9,
                                    textColor=NAVY, leading=12)
        st_cell = ParagraphStyle('cell', fontName='Helvetica', fontSize=8, leading=11,
                                 textColor=NAVY)
        st_cell_r = ParagraphStyle('cellr', parent=st_cell, alignment=TA_RIGHT)
        st_cell_b = ParagraphStyle('cellb', parent=st_cell, fontName='Helvetica-Bold')
        st_cell_b_r = ParagraphStyle('cellbr', parent=st_cell_b, alignment=TA_RIGHT)
        st_hdr = ParagraphStyle('hdr', parent=st_cell_b, textColor=WHITE)
        st_hdr_r = ParagraphStyle('hdrr', parent=st_hdr, alignment=TA_RIGHT)
        st_pill = ParagraphStyle('pill', fontName='Helvetica-Bold', fontSize=7.5,
                                 textColor=WHITE, alignment=TA_CENTER)
        st_note = ParagraphStyle('note', fontName='Helvetica-Oblique', fontSize=7.5,
                                 textColor=SLATE, leading=10)

        def section_heading(text, width):
            bar = Table([['']], colWidths=[1.2 * mm], rowHeights=[3.6 * mm])
            bar.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), TEAL)]))
            t = Table([[bar, Paragraph(text, st_section)]], colWidths=[3 * mm, width - 3 * mm])
            t.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ]))
            return t

        def status_pill(status, width):
            pill = Table([[Paragraph(STATUS_LABELS.get(status, status), st_pill)]],
                         colWidths=[width])
            pill.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), STATUS_COLORS.get(status, SLATE)),
                ('TOPPADDING', (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ]))
            return pill

        story = []

        # ── Header ──────────────────────────────────────────────────────────
        logo = RLImage(logo_path, width=15 * mm, height=15 * mm, kind='proportional')
        tenant = getattr(request, 'tenant', None)
        header_subs = []
        if tenant:
            addr = ', '.join(x for x in [getattr(tenant, 'address', '') or '',
                                         getattr(tenant, 'city', '') or ''] if x)
            if addr:
                header_subs.append(addr)
            contact = ' · '.join(x for x in [getattr(tenant, 'phone', '') or '',
                                             getattr(tenant, 'email', '') or ''] if x)
            if contact:
                header_subs.append(contact)
        left_block = [Paragraph(business_name, st_brand)] + \
            [Paragraph(s, st_brand_sub) for s in header_subs]
        date_range = ''
        if request.query_params.get('date_from') or request.query_params.get('date_to'):
            date_range = (f"Date range: {request.query_params.get('date_from') or '—'}"
                          f" to {request.query_params.get('date_to') or '—'}")
        right_block = [
            Paragraph('OPERATING EXPENSES REPORT', st_doctitle),
            Paragraph(f'Generated {_tz.now().strftime("%d %b %Y, %H:%M")}', st_docsub),
            *([Paragraph(date_range, st_docsub)] if date_range else []),
        ]
        header_tbl = Table([[logo, left_block, right_block]],
                           colWidths=[17 * mm, PAGE_W - 17 * mm - 80 * mm, 80 * mm])
        header_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, 0), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(header_tbl)
        story.append(Spacer(1, 4 * mm))
        strip = Table([['']], colWidths=[PAGE_W], rowHeights=[1.6 * mm])
        strip.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), TEAL)]))
        story.append(strip)
        story.append(Spacer(1, 5 * mm))

        # ── KPI cards ──────────────────────────────────────────────────────
        def kpi_card(label, value, sub=''):
            rows = [[Paragraph(label, st_label)], [Paragraph(value, st_value)]]
            if sub:
                rows.append([Paragraph(sub, st_value_sub)])
            return rows

        kpis = [
            kpi_card('TOTAL EXPENSES', money(total_amount), f'{len(expenses)} records'),
            kpi_card('THIS MONTH', money(this_month), 'amount'),
            kpi_card('PENDING', money(pending_value), f"{status_agg.get('pending', {}).get('count', 0)} awaiting approval"),
            kpi_card('PAID', money(paid_value), f"{status_agg.get('paid', {}).get('count', 0)} settled"),
            kpi_card('TAX', money(total_tax), 'across all expenses'),
        ]
        gap = 4 * mm
        kpi_w = (PAGE_W - 4 * gap) / 5
        kpi_tbl = Table(
            [[kpis[0], '', kpis[1], '', kpis[2], '', kpis[3], '', kpis[4]]],
            colWidths=[kpi_w, gap, kpi_w, gap, kpi_w, gap, kpi_w, gap, kpi_w],
        )
        kpi_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BACKGROUND', (0, 0), (0, 0), BG), ('BOX', (0, 0), (0, 0), 0.75, BORDER),
            ('BACKGROUND', (2, 0), (2, 0), BG), ('BOX', (2, 0), (2, 0), 0.75, BORDER),
            ('BACKGROUND', (4, 0), (4, 0), BG), ('BOX', (4, 0), (4, 0), 0.75, BORDER),
            ('BACKGROUND', (6, 0), (6, 0), BG), ('BOX', (6, 0), (6, 0), 0.75, BORDER),
            ('BACKGROUND', (8, 0), (8, 0), BG), ('BOX', (8, 0), (8, 0), 0.75, BORDER),
            ('LEFTPADDING', (0, 0), (-1, -1), 9),
            ('RIGHTPADDING', (0, 0), (-1, -1), 9),
            ('TOPPADDING', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        story.append(kpi_tbl)
        story.append(Spacer(1, 6 * mm))

        # ── Status breakdown + Top categories (side by side) ────────────────
        half = (PAGE_W - 10 * mm) / 2
        status_rows = [[
            Paragraph('STATUS', st_hdr), Paragraph('EXPENSES', st_hdr_r),
            Paragraph('VALUE (KSH)', st_hdr_r),
        ]]
        for key in ('pending', 'approved', 'paid', 'rejected', 'cancelled'):
            agg = status_agg.get(key)
            if not agg:
                continue
            status_rows.append([
                status_pill(key, 30 * mm),
                Paragraph(str(agg['count']), st_cell_r),
                Paragraph(f"{agg['value']:,.2f}", st_cell_r),
            ])
        status_tbl = Table(status_rows, colWidths=[half - 44 * mm, 22 * mm, 22 * mm])
        status_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        category_rows = [[
            Paragraph('CATEGORY', st_hdr),
            Paragraph('TOTAL SPEND (KSH)', st_hdr_r),
        ]]
        for name, value in top_categories:
            category_rows.append([
                Paragraph(name, st_cell),
                Paragraph(f'{value:,.2f}', st_cell_r),
            ])
        if not top_categories:
            category_rows.append([Paragraph('No data', st_note), ''])
        category_tbl = Table(category_rows, colWidths=[half - 34 * mm, 34 * mm])
        category_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        panels = Table([[
            [section_heading('STATUS BREAKDOWN', half), Spacer(1, 3 * mm), status_tbl],
            '',
            [section_heading('TOP CATEGORIES BY SPEND', half), Spacer(1, 3 * mm), category_tbl],
        ]], colWidths=[half, 10 * mm, half])
        panels.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(panels)
        story.append(Spacer(1, 6 * mm))

        # ── Expense ledger ──────────────────────────────────────────────────
        story.append(section_heading(f'EXPENSES ({len(expenses)})', PAGE_W))
        story.append(Spacer(1, 3.5 * mm))
        ledger_rows = [[
            Paragraph('#', st_hdr), Paragraph('REFERENCE', st_hdr),
            Paragraph('TITLE', st_hdr), Paragraph('CATEGORY', st_hdr),
            Paragraph('VENDOR', st_hdr), Paragraph('METHOD', st_hdr),
            Paragraph('DATE', st_hdr_r), Paragraph('STATUS', st_hdr),
            Paragraph('AMOUNT (KSH)', st_hdr_r),
        ]]
        for i, e in enumerate(expenses, 1):
            ledger_rows.append([
                Paragraph(str(i), st_cell),
                Paragraph(e.reference or '—', st_cell_b),
                Paragraph(e.title or '', st_cell),
                Paragraph(e.category.name if e.category else '—', st_cell),
                Paragraph(e.vendor or '—', st_cell),
                Paragraph(e.get_payment_method_display(), st_cell),
                Paragraph(fmt_date(e.expense_date), st_cell_r),
                status_pill(e.status, 26 * mm),
                Paragraph(f'{float(e.amount or 0):,.2f}', st_cell_b_r),
            ])
        if not expenses:
            ledger_rows.append([Paragraph('No expenses found for the selected filters.', st_note), '', '', '', '', '', '', '', ''])
        ledger_w = [8 * mm, 30 * mm, 52 * mm, 32 * mm, 40 * mm, 26 * mm, 26 * mm, 30 * mm, 32 * mm]
        ledger_tbl = Table(ledger_rows, colWidths=ledger_w, repeatRows=1)
        ledger_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 0), (-1, 0), 1, TEAL),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
        ]))
        story.append(ledger_tbl)

        # ── Footer (every page) ─────────────────────────────────────────────
        def draw_footer(canvas, doc_):
            canvas.saveState()
            canvas.setFillColor(NAVY)
            canvas.rect(0, PAGE[1] - 4 * mm, PAGE[0], 4 * mm, stroke=0, fill=1)
            canvas.setFillColor(TEAL)
            canvas.rect(0, PAGE[1] - 5.2 * mm, PAGE[0], 1.2 * mm, stroke=0, fill=1)
            canvas.setStrokeColor(BORDER)
            canvas.setLineWidth(0.5)
            canvas.line(15 * mm, 13 * mm, PAGE[0] - 15 * mm, 13 * mm)
            canvas.setFont('Helvetica', 7.5)
            canvas.setFillColor(SLATE)
            canvas.drawCentredString(PAGE[0] / 2, 8.5 * mm, 'Powered by AdhereMed')
            canvas.setFont('Helvetica-Bold', 8)
            canvas.setFillColor(TEAL)
            canvas.drawCentredString(PAGE[0] / 2, 4.5 * mm, 'www.adheremed.co')
            canvas.linkURL(
                'https://www.adheremed.co',
                (PAGE[0] / 2 - 60, 3.5 * mm, PAGE[0] / 2 + 60, 7.5 * mm),
            )
            canvas.restoreState()

        doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
        buf.seek(0)
        stamp = _tz.now().strftime('%Y%m%d')
        fname = f'operating_expenses_report_{stamp}.pdf'
        return HttpResponse(
            buf.read(),
            content_type='application/pdf',
            headers={'Content-Disposition': f'attachment; filename="{fname}"'},
        )

    @action(detail=False, methods=['get'])
    def summary(self, request):
        """Aggregate KPIs + breakdowns for dashboards."""
        qs = self.get_queryset()
        today = timezone.localdate()
        month_start = today.replace(day=1)

        totals = qs.aggregate(
            total=Sum('amount'),
            count=Count('id'),
        )
        by_status = list(qs.values('status').annotate(total=Sum('amount'), count=Count('id')))
        by_method = list(qs.values('payment_method').annotate(total=Sum('amount'), count=Count('id')))
        by_category = list(
            qs.values('category', 'category__name', 'category__color')
            .annotate(total=Sum('amount'), count=Count('id'))
            .order_by('-total')[:10]
        )
        this_month_total = qs.filter(expense_date__gte=month_start).aggregate(s=Sum('amount'))['s'] or 0
        pending_total = qs.filter(status=Expense.Status.PENDING).aggregate(s=Sum('amount'))['s'] or 0
        paid_total = qs.filter(status=Expense.Status.PAID).aggregate(s=Sum('amount'))['s'] or 0
        overdue = qs.filter(
            due_date__isnull=False, due_date__lt=today,
        ).exclude(status__in=[Expense.Status.PAID, Expense.Status.CANCELLED, Expense.Status.REJECTED])
        overdue_total = overdue.aggregate(s=Sum('amount'))['s'] or 0

        # last 12 months trend
        trend = []
        y, m = month_start.year, month_start.month
        for i in range(11, -1, -1):
            mm = m - i
            yy = y
            while mm <= 0:
                mm += 12
                yy -= 1
            start = month_start.replace(year=yy, month=mm, day=1)
            nm = mm + 1
            ny = yy
            if nm > 12:
                nm = 1
                ny += 1
            end = start.replace(year=ny, month=nm, day=1)
            total = qs.filter(expense_date__gte=start, expense_date__lt=end).aggregate(s=Sum('amount'))['s'] or 0
            trend.append({'month': start.isoformat(), 'total': float(total)})

        return Response({
            'total': float(totals['total'] or 0),
            'count': totals['count'] or 0,
            'this_month_total': float(this_month_total),
            'pending_total': float(pending_total),
            'paid_total': float(paid_total),
            'overdue_total': float(overdue_total),
            'overdue_count': overdue.count(),
            'by_status': [
                {'status': r['status'], 'total': float(r['total'] or 0), 'count': r['count']}
                for r in by_status
            ],
            'by_method': [
                {'method': r['payment_method'], 'total': float(r['total'] or 0), 'count': r['count']}
                for r in by_method
            ],
            'by_category': [
                {
                    'category_id': r['category'],
                    'name': r['category__name'] or 'Uncategorized',
                    'color': r['category__color'] or '',
                    'total': float(r['total'] or 0),
                    'count': r['count'],
                }
                for r in by_category
            ],
            'trend': trend,
        })

    # ── Excel import/export (template, review, commit) ────────────────────────
    _IMPORT_COLUMNS = {
        'reference': ['Reference', 'reference'],
        'title': ['Title', 'title'],
        'category': ['Category', 'category'],
        'vendor': ['Vendor', 'vendor'],
        'amount': ['Amount (KSh)', 'amount'],
        'tax_amount': ['Tax (KSh)', 'tax_amount'],
        'expense_date': ['Expense Date', 'expense_date'],
        'due_date': ['Due Date', 'due_date'],
        'payment_method': ['Payment Method', 'payment_method'],
        'payment_reference': ['Payment Reference', 'payment_reference'],
        'status': ['Status', 'status'],
        'is_recurring': ['Recurring', 'is_recurring'],
        'recurring_period': ['Recurring Period', 'recurring_period'],
        'notes': ['Notes', 'notes'],
    }
    _PAYMENT_METHODS = {'cash', 'mpesa', 'bank', 'card', 'cheque', 'other'}
    _STATUSES = {'pending', 'approved', 'paid', 'rejected', 'cancelled'}
    _RECUR_PERIODS = {'daily', 'weekly', 'monthly', 'quarterly', 'yearly'}

    _EXPORT_HEADERS = [
        'Reference', 'Title', 'Category', 'Vendor', 'Amount (KSh)', 'Tax (KSh)',
        'Total (KSh)', 'Expense Date', 'Due Date', 'Payment Method',
        'Payment Reference', 'Status', 'Recurring', 'Recurring Period',
        'Submitted By', 'Notes', 'Created At',
    ]

    def _expense_row(self, e):
        return [
            e.reference, e.title, e.category.name if e.category else '', e.vendor or '',
            float(e.amount), float(e.tax_amount or 0), float(e.total_amount or 0),
            e.expense_date.isoformat(), e.due_date.isoformat() if e.due_date else '',
            e.payment_method, e.payment_reference or '', e.status,
            'Yes' if e.is_recurring else 'No', e.recurring_period or '',
            str(e.submitted_by) if e.submitted_by else '', e.notes or '',
            e.created_at.strftime('%Y-%m-%d'),
        ]

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export expenses as .xlsx or .csv (respects current filters)."""
        # NB: use ?fmt= — DRF reserves the `format` query param.
        fmt = request.query_params.get('fmt', 'csv').lower()
        qs = self.filter_queryset(self.get_queryset())

        if fmt == 'excel':
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment
                from openpyxl.utils import get_column_letter
            except ImportError:
                return Response({'detail': 'openpyxl not installed.'}, status=500)

            wb = Workbook()
            ws = wb.active
            ws.title = 'Expenses'
            header_font = Font(bold=True, color='FFFFFF', size=11)
            header_fill = PatternFill(fill_type='solid', fgColor='0EA5E9')
            for col_idx, header in enumerate(self._EXPORT_HEADERS, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
                ws.column_dimensions[get_column_letter(col_idx)].width = \
                    max(12, min(len(header) + 4, 32))
            for row_idx, e in enumerate(qs, 2):
                for col_idx, val in enumerate(self._expense_row(e), 1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
            ws.freeze_panes = 'A2'

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            fname = f'expenses_{timezone.now():%Y%m%d}.xlsx'
            return HttpResponse(
                buf.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                headers={'Content-Disposition': f'attachment; filename="{fname}"'},
            )

        fname = f'expenses_{timezone.now():%Y%m%d}.csv'
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{fname}"'
        writer = csv.writer(response)
        writer.writerow(self._EXPORT_HEADERS)
        for e in qs:
            writer.writerow(self._expense_row(e))
        return response

    @staticmethod
    def _parse_import_date(v):
        """Accept date/datetime objects or common date string formats."""
        if v in (None, ''):
            return None
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        s = str(v).strip()
        try:
            return datetime.fromisoformat(s.split()[0]).date()
        except ValueError:
            pass
        for f in ('%Y-%m-%d %H:%M:%S', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y',
                  '%Y/%m/%d', '%d %b %Y', '%b %d, %Y'):
            try:
                return datetime.strptime(s, f).date()
            except ValueError:
                continue
        raise ValueError(f'"{s}" is not a valid date')

    def _normalize_import_rows(self, uploaded):
        """Parse an uploaded .xlsx/.csv into (row_num, {field: raw_value}) pairs."""
        filename = (uploaded.name or '').lower()
        if filename.endswith('.csv'):
            text = uploaded.read().decode('utf-8-sig')
            rows = [r for r in csv.reader(io.StringIO(text))]
        elif filename.endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(uploaded.read()), read_only=True, data_only=True)
                ws = wb.active
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
            except Exception:
                raise ValueError('Could not read the spreadsheet. Save it as .xlsx or .csv and try again.')
        else:
            raise ValueError('Unsupported file type. Use .xlsx or .csv')

        if not rows:
            raise ValueError('The file is empty.')
        headers = [str(h).strip() if h is not None else '' for h in rows[0]]
        col = {h: i for i, h in enumerate(headers)}
        if 'Title' not in col and 'title' not in col:
            raise ValueError('Missing the "Title" column. Use the import template.')

        def get(row, aliases):
            for a in aliases:
                idx = col.get(a)
                if idx is not None and idx < len(row):
                    v = row[idx]
                    if v is not None and str(v).strip() != '':
                        return v
            return None

        parsed = []
        for num, row in enumerate(rows[1:], 2):
            if not row or all(v is None or str(v).strip() == '' for v in row):
                continue
            parsed.append((num, {f: get(row, aliases)
                                 for f, aliases in self._IMPORT_COLUMNS.items()}))
        return parsed

    def _clean_import_row(self, raw, is_new):
        """Validate one raw expense row. Returns (cleaned, errors)."""
        cleaned, errors = {}, {}

        title = str(raw.get('title') or '').strip()
        if not title:
            errors['title'] = 'Required'
        cleaned['title'] = title or None

        cleaned['reference'] = str(raw.get('reference') or '').strip() or None

        # Amount: required for new expenses, optional (keep existing) on update
        v = raw.get('amount')
        if v in (None, ''):
            cleaned['amount'] = None
            if is_new:
                errors['amount'] = 'Required'
        else:
            try:
                n = float(v)
                if n < 0:
                    raise ValueError
                cleaned['amount'] = n
            except (TypeError, ValueError):
                cleaned['amount'] = str(v)
                errors['amount'] = 'Must be a number ≥ 0'

        v = raw.get('tax_amount')
        if v in (None, ''):
            cleaned['tax_amount'] = None
        else:
            try:
                n = float(v)
                if n < 0:
                    raise ValueError
                cleaned['tax_amount'] = n
            except (TypeError, ValueError):
                cleaned['tax_amount'] = str(v)
                errors['tax_amount'] = 'Must be a number ≥ 0'

        for f in ('expense_date', 'due_date'):
            v = raw.get(f)
            if v in (None, ''):
                cleaned[f] = None
                if f == 'expense_date' and is_new:
                    errors[f] = 'Required'
                continue
            try:
                cleaned[f] = self._parse_import_date(v).isoformat()
            except ValueError as exc:
                cleaned[f] = str(v)
                errors[f] = str(exc)

        def _enum(field, valid, default=None):
            v = raw.get(field)
            if v in (None, ''):
                cleaned[field] = default
                return
            s = str(v).strip().lower()
            if s in valid:
                cleaned[field] = s
            else:
                cleaned[field] = str(v)
                errors[field] = 'Invalid value'

        _enum('payment_method', self._PAYMENT_METHODS, 'cash')
        _enum('status', self._STATUSES, 'pending')
        _enum('recurring_period', self._RECUR_PERIODS)

        v = raw.get('is_recurring')
        if v in (None, ''):
            cleaned['is_recurring'] = None
        else:
            s = str(v).strip().lower()
            if s in ('yes', 'true', '1', 'no', 'false', '0'):
                cleaned['is_recurring'] = 'Yes' if s in ('yes', 'true', '1') else 'No'
            else:
                cleaned['is_recurring'] = str(v)
                errors['is_recurring'] = 'Must be Yes or No'

        for f in ('category', 'vendor', 'payment_reference', 'notes'):
            v = raw.get(f)
            cleaned[f] = str(v).strip() if v not in (None, '') else None
        return cleaned, errors

    @action(detail=False, methods=['get'], url_path='import-template')
    def import_template(self, request):
        """Styled .xlsx import template with example row, dropdowns & instructions."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.worksheet.datavalidation import DataValidation
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response({'detail': 'openpyxl not installed.'}, status=500)

        HEADERS = [
            ('Reference', 16, 'blank = auto-generated'),
            ('Title', 34, 'required'),
            ('Category', 22, 'auto-created if missing'),
            ('Vendor', 24, 'payee / supplier name'),
            ('Amount (KSh)', 15, 'required · number ≥ 0'),
            ('Tax (KSh)', 12, 'number ≥ 0'),
            ('Expense Date', 15, 'required · YYYY-MM-DD'),
            ('Due Date', 14, 'optional · YYYY-MM-DD'),
            ('Payment Method', 16, 'cash / mpesa / bank / card / cheque / other'),
            ('Payment Reference', 18, 'mpesa code, cheque #, etc.'),
            ('Status', 12, 'pending / approved / paid / rejected / cancelled'),
            ('Recurring', 11, 'Yes / No'),
            ('Recurring Period', 17, 'daily / weekly / monthly / quarterly / yearly'),
            ('Notes', 32, 'free text'),
        ]
        EXAMPLE = ['', 'Rent - September', 'Facilities', 'Mombasa Properties Ltd', 85000, 0,
                   '2026-09-01', '', 'bank', 'TRF-88291', 'approved', 'No', '',
                   'Monthly office rent']

        wb = Workbook()
        ws = wb.active
        ws.title = 'Expenses'
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(fill_type='solid', fgColor='0EA5E9')
        thin = Side(style='thin', color='E2E8F0')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for c, (title, width, _) in enumerate(HEADERS, 1):
            cell = ws.cell(row=1, column=c, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[1].height = 24
        ws.freeze_panes = 'A2'

        for c, val in enumerate(EXAMPLE, 1):
            cell = ws.cell(row=2, column=c, value=val)
            cell.font = Font(italic=True, color='64748B')
            cell.border = border

        def _dropdown(formula, col_letter, error_msg):
            dv = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
            dv.error = error_msg
            ws.add_data_validation(dv)
            dv.add(f'{col_letter}3:{col_letter}1000')

        _dropdown('"cash,mpesa,bank,card,cheque,other"', 'I', 'Invalid payment method')
        _dropdown('"pending,approved,paid,rejected,cancelled"', 'K', 'Invalid status')
        _dropdown('"Yes,No"', 'L', 'Must be Yes or No')
        _dropdown('"daily,weekly,monthly,quarterly,yearly"', 'M', 'Invalid period')

        ins = wb.create_sheet('Instructions')
        ins.column_dimensions['A'].width = 6
        ins.column_dimensions['B'].width = 96
        title_cell = ins.cell(row=1, column=2, value='AdhereMed · Expenses Excel Import')
        title_cell.font = Font(bold=True, size=14, color='0F172A')
        lines = [
            '',
            'HOW TO USE',
            '1. Fill one row per expense in the "Expenses" sheet (or edit an exported file).',
            '2. "Title", "Amount" and "Expense Date" are required for new expenses.',
            '3. A row with an existing "Reference" updates that expense. A blank reference creates a new one.',
            '4. "Category" is matched by name and auto-created if missing.',
            '5. Blank cells keep existing values on update; defaults apply to new expenses.',
            '6. Supported formats: .xlsx and .csv',
            '',
            'REVIEW WORKFLOW',
            '· Upload the file on the Excel Import / Export page.',
            '· Every row is checked and shown in a review table where you can edit or remove rows.',
            '· Only when you press "Submit import" is anything saved.',
        ]
        for i, line in enumerate(lines, 3):
            cell = ins.cell(row=i, column=2, value=line)
            if line.isupper():
                cell.font = Font(bold=True, size=11, color='0EA5E9')
            else:
                cell.font = Font(size=10, color='334155')

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f'expenses_import_template_{timezone.now():%Y%m%d}.xlsx'
        return HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{fname}"'},
        )

    @action(detail=False, methods=['post'], url_path='import-preview')
    def import_preview(self, request):
        """Parse an uploaded spreadsheet and return the rows for review."""
        uploaded = request.FILES.get('file')
        if not uploaded:
            return Response({'detail': 'No file provided.'}, status=400)
        try:
            parsed = self._normalize_import_rows(uploaded)
        except ValueError as exc:
            return Response({'detail': str(exc)}, status=400)

        existing = {
            e.reference.lower(): {'id': e.id}
            for e in Expense.objects.only('id', 'reference')
            if e.reference
        }

        rows, seen = [], {}
        counts = {'new': 0, 'update': 0, 'error': 0}
        for num, raw in parsed:
            ref = str(raw.get('reference') or '').strip()
            key = ref.lower() if ref else None
            ex = existing.get(key) if key else None
            cleaned, errors = self._clean_import_row(raw, is_new=ex is None)
            if key and key in seen:
                errors['reference'] = f'Duplicate of row {seen[key]}'
            elif key:
                seen[key] = num
            status_ = 'error' if errors else ('update' if ex else 'new')
            counts[status_] += 1
            rows.append({
                'row_num': num,
                **cleaned,
                'row_status': status_,
                'existing_id': ex['id'] if ex else None,
                'errors': errors,
            })
        return Response({
            'rows': rows,
            'summary': {
                'total': len(rows),
                'new': counts['new'],
                'update': counts['update'],
                'error': counts['error'],
            },
        })

    @action(detail=False, methods=['post'], url_path='import-commit')
    def import_commit(self, request):
        """Apply reviewed (and possibly edited) expense rows."""
        items = request.data.get('items')
        if not isinstance(items, list) or not items:
            return Response({'detail': 'items must be a non-empty list.'}, status=400)

        created = updated = failed = 0
        results = []

        def _num(item, field, default=None):
            v = item.get(field)
            if v in (None, ''):
                return default
            return Decimal(str(v))

        for idx, item in enumerate(items, 1):
            row_label = item.get('row_num') or idx
            title = str(item.get('title') or '').strip()
            ref = str(item.get('reference') or '').strip() or None
            try:
                if not title:
                    raise ValueError('Title is required')

                expense = None
                if ref:
                    expense = Expense.objects.filter(reference__iexact=ref).first()
                is_new = expense is None
                if is_new:
                    expense = Expense(
                        reference=ref or f'EXP-{uuid.uuid4().hex[:8].upper()}',
                        title=title,
                        submitted_by=request.user if request.user.is_authenticated else None,
                    )
                else:
                    expense.title = title

                amount = _num(item, 'amount')
                if amount is not None:
                    if amount < 0:
                        raise InvalidOperation()
                    expense.amount = amount
                elif is_new:
                    raise ValueError('Amount is required')

                tax = _num(item, 'tax_amount')
                if tax is not None:
                    if tax < 0:
                        raise InvalidOperation()
                    expense.tax_amount = tax
                elif is_new:
                    expense.tax_amount = Decimal('0')

                for f in ('expense_date', 'due_date'):
                    v = item.get(f)
                    if v not in (None, ''):
                        expense.__setattr__(f, self._parse_import_date(v))
                    elif is_new and f == 'expense_date':
                        raise ValueError('Expense date is required')

                if item.get('category'):
                    cat, _ = ExpenseCategory.objects.get_or_create(
                        name=str(item['category']).strip())
                    expense.category = cat

                pm = str(item.get('payment_method') or '').strip().lower()
                if pm in self._PAYMENT_METHODS:
                    expense.payment_method = pm
                elif is_new:
                    expense.payment_method = 'cash'

                st = str(item.get('status') or '').strip().lower()
                if st in self._STATUSES:
                    expense.status = st
                    if st == Expense.Status.APPROVED and not expense.approved_at:
                        expense.approved_at = timezone.now()
                        if request.user.is_authenticated:
                            expense.approved_by = request.user
                    if st == Expense.Status.PAID and not expense.paid_at:
                        expense.paid_at = timezone.now()
                elif is_new:
                    expense.status = 'pending'

                rec = item.get('is_recurring')
                if isinstance(rec, str):
                    expense.is_recurring = rec.strip().lower() in ('yes', 'true', '1')
                elif isinstance(rec, bool):
                    expense.is_recurring = rec

                rp = str(item.get('recurring_period') or '').strip().lower()
                if rp in self._RECUR_PERIODS:
                    expense.recurring_period = rp

                for f in ('vendor', 'payment_reference', 'notes'):
                    v = item.get(f)
                    if v not in (None, ''):
                        setattr(expense, f, str(v).strip())

                expense.save()
                if is_new:
                    created += 1
                else:
                    updated += 1
                results.append({'row': row_label, 'name': title,
                                 'status': 'created' if is_new else 'updated',
                                 'reference': expense.reference})
            except Exception as exc:
                failed += 1
                results.append({'row': row_label, 'name': title or '(no title)',
                                'status': 'error', 'error': str(exc) or 'invalid row'})

        return Response({'created': created, 'updated': updated,
                         'failed': failed, 'results': results})
