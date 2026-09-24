import csv
import io
import os
from datetime import datetime

from django.http import HttpResponse
from rest_framework import viewsets, filters, status
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend

from config.branch_scope import BranchScopedMixin
from .models import PurchaseOrder, GoodsReceivedNote
from .serializers import (
    PurchaseOrderSerializer,
    GoodsReceivedNoteSerializer,
    revert_received_items,
)


class PurchaseOrderViewSet(BranchScopedMixin, viewsets.ModelViewSet):
    queryset = PurchaseOrder.objects.select_related('supplier', 'ordered_by', 'branch').prefetch_related('grns').all()
    serializer_class = PurchaseOrderSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['status', 'supplier', 'branch']
    search_fields = ['po_number', 'supplier__name']
    ordering_fields = ['created_at', 'order_date', 'total_cost']

    def get_queryset(self):
        """Extend the base queryset with optional date range filters
        (date_from / date_to on order_date) for list, export and reports."""
        qs = super().get_queryset()
        params = self.request.query_params
        date_from = params.get('date_from')
        date_to = params.get('date_to')
        if date_from:
            qs = qs.filter(order_date__gte=date_from)
        if date_to:
            qs = qs.filter(order_date__lte=date_to)
        return qs

    @action(detail=True, methods=['get'])
    def return_preview(self, request, pk=None):
        """Show what would happen if this PO were returned (without applying)."""
        po = self.get_object()
        preview = []
        from inventory.models import MedicationStock, StockBatch
        for it in (po.items or []):
            if not it.get('_synced') or it.get('_returned'):
                continue
            try:
                stock = MedicationStock.objects.get(pk=it['medication_stock_id'])
            except MedicationStock.DoesNotExist:
                continue
            batch = StockBatch.objects.filter(stock=stock, batch_number=it.get('batch_number')).first()
            received = int(it.get('qty') or 0)
            consumed = 0
            remaining = 0
            if batch:
                remaining = int(batch.quantity_remaining)
                consumed = max(0, int(batch.quantity_received) - remaining)
            preview.append({
                'name': it.get('name') or stock.medication_name,
                'received': received,
                'consumed': consumed,
                'remaining': remaining,
                'current_cost_price': float(stock.cost_price or 0),
                'previous_cost_price': it.get('_prev_cost_price'),
                'current_selling_price': float(stock.selling_price or 0),
                'previous_selling_price': it.get('_prev_selling_price'),
            })
        return Response({
            'po_number': po.po_number,
            'status': po.status,
            'items': preview,
            'has_consumed': any(p['consumed'] > 0 for p in preview),
        })

    @action(detail=True, methods=['post'])
    def return_purchase(self, request, pk=None):
        """Return a previously received PO. Reverses stock and prices.

        Pass `{"force": true}` to allow reverting even if some stock was already used.
        """
        po = self.get_object()
        if po.status not in (PurchaseOrder.Status.RECEIVED, PurchaseOrder.Status.PARTIAL):
            return Response(
                {'detail': f'Only received purchase orders can be returned (current status: {po.status}).'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        force = bool(request.data.get('force'))
        try:
            result = revert_received_items(po, force=force)
        except ValueError as exc:
            payload = exc.args[0] if exc.args else {'message': str(exc)}
            if isinstance(payload, dict):
                payload['needs_force'] = True
                return Response(payload, status=status.HTTP_409_CONFLICT)
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        po.refresh_from_db()
        data = self.get_serializer(po).data
        return Response({**result, 'purchase_order': data})

    # ── Excel import for PO items ────────────────────────────────────────────
    _ITEM_IMPORT_COLUMNS = {
        'name': ['Item Name', 'Medication Name', 'item', 'name', 'medication_name'],
        'qty': ['Quantity', 'Qty', 'quantity', 'qty'],
        'unit_cost': ['Unit Cost (KSh)', 'Unit Cost', 'unit_cost', 'cost_price', 'cost'],
        'unit_selling_price': ['Unit Selling Price (KSh)', 'Unit Selling Price',
                               'Selling Price (KSh)', 'selling_price', 'unit_selling_price'],
        'discount_percent': ['Discount %', 'Disc %', 'discount_percent'],
        'tax_percent': ['VAT %', 'Tax %', 'tax_percent'],
        'batch_number': ['Batch #', 'Batch Number', 'batch_number', 'batch'],
        'expiry_date': ['Expiry Date', 'Expiry', 'expiry_date'],
    }

    def _parse_item_spreadsheet(self, uploaded):
        """Parse an uploaded .xlsx/.csv into (row_num, {field: raw_value}) pairs."""
        filename = (uploaded.name or '').lower()
        if filename.endswith('.csv'):
            text = uploaded.read().decode('utf-8-sig')
            rows = [list(r) for r in csv.reader(io.StringIO(text))]
        elif filename.endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(uploaded.read()), read_only=True, data_only=True)
                rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
            except Exception:
                raise ValueError('Could not read the spreadsheet. Save it as .xlsx or .csv and try again.')
        else:
            raise ValueError('Unsupported file type. Use .xlsx or .csv')

        if not rows:
            raise ValueError('The file is empty.')
        headers = [str(h).strip() if h is not None else '' for h in rows[0]]
        col = {h: i for i, h in enumerate(headers)}
        if not any(a in col for a in self._ITEM_IMPORT_COLUMNS['name']):
            raise ValueError('Missing the "Item Name" column. Use the import template.')

        def get(row, key):
            for a in self._ITEM_IMPORT_COLUMNS[key]:
                idx = col.get(a)
                if idx is not None and idx < len(row):
                    v = row[idx]
                    if v is not None and str(v).strip() != '':
                        return v
            return None

        parsed = []
        for num, row in enumerate(rows[1:], 2):
            if not row or all(v is None or str(v).strip() == '' for v in row):
                continue
            parsed.append((num, {k: get(row, k) for k in self._ITEM_IMPORT_COLUMNS}))
        if not parsed:
            raise ValueError('No data rows found below the header row.')
        return parsed

    @staticmethod
    def _clean_number(raw, default=0.0):
        if raw is None or str(raw).strip() == '':
            return default, None
        try:
            return float(str(raw).replace(',', '')), None
        except ValueError:
            return default, 'Invalid number'

    @staticmethod
    def _clean_date(raw):
        if raw is None or str(raw).strip() == '':
            return ''
        if hasattr(raw, 'strftime'):
            return raw.strftime('%Y-%m-%d')
        s = str(raw).strip()
        for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%d.%m.%Y'):
            try:
                return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
            except ValueError:
                continue
        return s

    def _resolve_branding(self, request):
        """Return (logo_path, business_name) for PDF documents."""
        from config.pdf_branding import resolve_branding
        return resolve_branding(request)

    @action(detail=True, methods=['post'], url_path='upload-proof',
            parser_classes=[MultiPartParser, FormParser])
    def upload_proof(self, request, pk=None):
        """Attach an optional proof image (signed note, receipt photo) to a PO."""
        po = self.get_object()
        image = request.FILES.get('image') or request.FILES.get('proof')
        if not image:
            return Response({'detail': 'No image file provided.'}, status=400)
        if not (image.content_type or '').startswith('image/'):
            return Response({'detail': 'File must be an image.'}, status=400)
        if po.proof_image:
            po.proof_image.delete(save=False)
        po.proof_image = image
        po.save(update_fields=['proof_image'])
        return Response({
            'proof_image_url': request.build_absolute_uri(po.proof_image.url),
        })

    @action(detail=True, methods=['post'], url_path='clear-proof')
    def clear_proof(self, request, pk=None):
        """Remove the proof image from a PO."""
        po = self.get_object()
        if po.proof_image:
            po.proof_image.delete(save=False)
            po.proof_image = None
            po.save(update_fields=['proof_image'])
        return Response({'proof_image_url': None})

    @action(detail=True, methods=['get'], url_path='po-pdf')
    def po_pdf(self, request, pk=None):
        """Generate a branded premium PDF of the Purchase Order (PO)."""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib.colors import HexColor
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.enums import TA_RIGHT, TA_CENTER
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
            )
        except ImportError:
            return Response({'detail': 'reportlab not installed.'}, status=500)

        po = self.get_object()
        tenant = getattr(request, 'tenant', None)
        logo_path, business_name = self._resolve_branding(request)

        NAVY = HexColor('#0F172A')
        TEAL = HexColor('#0D9488')
        SLATE = HexColor('#64748B')
        BORDER = HexColor('#E2E8F0')
        BG = HexColor('#F8FAFC')
        WHITE = HexColor('#FFFFFF')
        STATUS_COLORS = {
            'draft': HexColor('#64748B'),
            'sent': HexColor('#2563EB'),
            'received': HexColor('#16A34A'),
            'partial': HexColor('#D97706'),
            'cancelled': HexColor('#DC2626'),
            'returned': HexColor('#7C3AED'),
        }
        status_color = STATUS_COLORS.get(po.status, SLATE)

        def money(v):
            try:
                return f'KSh {float(v or 0):,.2f}'
            except (TypeError, ValueError):
                return 'KSh 0.00'

        def fmt_date(d):
            return d.strftime('%d %b %Y') if d else '—'

        PAGE_W = A4[0] - 30 * mm  # usable width

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            leftMargin=15 * mm, rightMargin=15 * mm,
            topMargin=12 * mm, bottomMargin=18 * mm,
            title=f'Purchase Order {po.po_number}',
        )

        # ── Styles ──────────────────────────────────────────────────────────
        st_brand = ParagraphStyle('brand', fontName='Helvetica-Bold', fontSize=13,
                                   textColor=NAVY, leading=16)
        st_brand_sub = ParagraphStyle('brandsub', fontName='Helvetica', fontSize=8,
                                      textColor=SLATE, leading=11)
        st_doctitle = ParagraphStyle('doctitle', fontName='Helvetica-Bold', fontSize=21,
                                     textColor=TEAL, leading=24, alignment=TA_RIGHT)
        st_docsub = ParagraphStyle('docsub', fontName='Helvetica', fontSize=9,
                                   textColor=SLATE, leading=12, alignment=TA_RIGHT)
        st_label = ParagraphStyle('label', fontName='Helvetica-Bold', fontSize=6.5,
                                  textColor=SLATE, leading=9)
        st_value = ParagraphStyle('value', fontName='Helvetica-Bold', fontSize=10,
                                  textColor=NAVY, leading=13)
        st_value_c = ParagraphStyle('valuec', parent=st_value, alignment=TA_CENTER)
        st_section = ParagraphStyle('section', fontName='Helvetica-Bold', fontSize=9,
                                    textColor=NAVY, leading=12)
        st_body = ParagraphStyle('body', fontName='Helvetica', fontSize=8.5,
                                 textColor=SLATE, leading=12)
        st_cell = ParagraphStyle('cell', fontName='Helvetica', fontSize=8.5, leading=12,
                                 textColor=NAVY)
        st_cell_r = ParagraphStyle('cellr', parent=st_cell, alignment=TA_RIGHT)
        st_cell_b = ParagraphStyle('cellb', parent=st_cell, fontName='Helvetica-Bold')
        st_cell_b_r = ParagraphStyle('cellbr', parent=st_cell_b, alignment=TA_RIGHT)
        st_hdr = ParagraphStyle('hdr', parent=st_cell_b, textColor=WHITE, fontSize=8)
        st_hdr_r = ParagraphStyle('hdrr', parent=st_hdr, alignment=TA_RIGHT)
        st_note = ParagraphStyle('note', fontName='Helvetica-Oblique', fontSize=7.5,
                                 textColor=SLATE, leading=10)
        st_pill = ParagraphStyle('pill', fontName='Helvetica-Bold', fontSize=8,
                                 textColor=WHITE, alignment=TA_CENTER)
        st_tot_lbl = ParagraphStyle('totlbl', fontName='Helvetica', fontSize=9,
                                    textColor=SLATE, leading=13)
        st_tot_lbl_b = ParagraphStyle('totlblb', parent=st_tot_lbl,
                                      fontName='Helvetica-Bold', textColor=WHITE)
        st_tot_val = ParagraphStyle('totval', fontName='Helvetica', fontSize=9,
                                    textColor=NAVY, leading=13, alignment=TA_RIGHT)
        st_tot_val_b = ParagraphStyle('totvalb', fontName='Helvetica-Bold', fontSize=12,
                                      textColor=WHITE, leading=15, alignment=TA_RIGHT)

        def section_heading(text):
            bar = Table([['']], colWidths=[1.2 * mm], rowHeights=[3.6 * mm])
            bar.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), TEAL)]))
            t = Table([[bar, Paragraph(text, st_section)]], colWidths=[3 * mm, PAGE_W - 3 * mm])
            t.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ]))
            return t

        story = []

        # ── Header: logo + tenant (left) · document title (right) ───────────
        logo = RLImage(logo_path, width=15 * mm, height=15 * mm, kind='proportional')
        tenant_subs = []
        if tenant:
            addr = ', '.join(x for x in [getattr(tenant, 'address', '') or '',
                                         getattr(tenant, 'city', '') or ''] if x)
            if addr:
                tenant_subs.append(addr)
            contact = ' · '.join(x for x in [getattr(tenant, 'phone', '') or '',
                                             getattr(tenant, 'email', '') or ''] if x)
            if contact:
                tenant_subs.append(contact)
        left_block = [Paragraph(business_name, st_brand)] + \
            [Paragraph(s, st_brand_sub) for s in tenant_subs]
        right_block = [
            Paragraph('PURCHASE ORDER', st_doctitle),
            Paragraph(f'PO · {po.po_number}', st_docsub),
        ]
        header_tbl = Table([[logo, left_block, right_block]],
                           colWidths=[17 * mm, PAGE_W - 17 * mm - 64 * mm, 64 * mm])
        header_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, 0), 8),
            ('RIGHTPADDING', (1, 0), (1, 0), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(header_tbl)
        story.append(Spacer(1, 4 * mm))

        # Brand strip
        strip = Table([['']], colWidths=[PAGE_W], rowHeights=[1.6 * mm])
        strip.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), TEAL)]))
        story.append(strip)
        story.append(Spacer(1, 5 * mm))

        # ── Meta cards row ──────────────────────────────────────────────────
        def meta_card(label, value_flowable):
            t = Table([[label], [value_flowable]], colWidths=[42 * mm])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), BG),
                ('BOX', (0, 0), (-1, -1), 0.75, BORDER),
                ('LEFTPADDING', (0, 0), (-1, -1), 9),
                ('RIGHTPADDING', (0, 0), (-1, -1), 9),
                ('TOPPADDING', (0, 0), (-1, 0), 7),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 1),
                ('TOPPADDING', (0, 1), (-1, 1), 1),
                ('BOTTOMPADDING', (0, 1), (-1, 1), 7),
            ]))
            return t

        pill = Table([[Paragraph(po.get_status_display(), st_pill)]], colWidths=[24 * mm])
        pill.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), status_color),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))

        gap = 3.75 * mm
        card_w = (PAGE_W - 3 * gap) / 4
        status_card = Table(
            [[Paragraph('STATUS', st_label)], [pill]],
            colWidths=[card_w],
            style=TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), BG),
                ('BOX', (0, 0), (-1, -1), 0.75, BORDER),
                ('LEFTPADDING', (0, 0), (-1, -1), 9),
                ('RIGHTPADDING', (0, 0), (-1, -1), 9),
                ('TOPPADDING', (0, 0), (-1, 0), 7),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 1),
                ('TOPPADDING', (0, 1), (-1, 1), 2),
                ('BOTTOMPADDING', (0, 1), (-1, 1), 6),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ]),
        )
        meta_row = Table([[
            meta_card('PO NUMBER', Paragraph(po.po_number, st_value)),
            '',
            meta_card('ORDER DATE', Paragraph(fmt_date(po.order_date), st_value)),
            '',
            meta_card('EXPECTED DELIVERY', Paragraph(fmt_date(po.expected_delivery), st_value)),
            '',
            status_card,
        ]], colWidths=[card_w, gap, card_w, gap, card_w, gap, card_w])
        meta_row.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(meta_row)
        story.append(Spacer(1, 7 * mm))

        # ── Supplier / Deliver To ───────────────────────────────────────────
        supplier = po.supplier
        supplier_block = []
        if supplier:
            supplier_block.append(Paragraph(supplier.name, ParagraphStyle(
                'sname', parent=st_cell_b, fontSize=10.5, leading=13)))
            for label, value in [
                ('Contact', getattr(supplier, 'contact_person', '') or ''),
                ('Phone', getattr(supplier, 'phone', '') or ''),
                ('Email', getattr(supplier, 'email', '') or ''),
                ('Address', (getattr(supplier, 'address', '') or '').replace('\n', ', ')),
                ('Payment terms', getattr(supplier, 'payment_terms', '') or ''),
            ]:
                if value:
                    supplier_block.append(Paragraph(f'{label}:  {value}', st_body))

        half = (PAGE_W - 8 * mm) / 2
        party_tbl = Table([[
            [section_heading('SUPPLIER'), Spacer(1, 3 * mm), *supplier_block],
            [section_heading('DELIVER TO'), Spacer(1, 3 * mm),
             Paragraph(po.branch.name if po.branch else business_name or '—',
                       ParagraphStyle('bname', parent=st_cell_b, fontSize=10.5, leading=13)),
             *( [Paragraph(f'Expected: {fmt_date(po.expected_delivery)}', st_body)]
                if po.expected_delivery else [] ),
             *( [Paragraph(f'Ordered by: {po.ordered_by.get_full_name() if hasattr(po.ordered_by, "get_full_name") else (po.ordered_by or "")}', st_body)]
                if po.ordered_by else [] )],
        ]], colWidths=[half, half])
        party_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (0, 0), 0),
            ('LEFTPADDING', (1, 0), (1, 0), 8 * mm),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(party_tbl)
        story.append(Spacer(1, 6 * mm))

        # ── Items table ──────────────────────────────────────────────────────
        story.append(section_heading('ORDER ITEMS'))
        story.append(Spacer(1, 3.5 * mm))
        item_rows = [[
            Paragraph('#', st_hdr), Paragraph('ITEM', st_hdr),
            Paragraph('QTY', st_hdr_r), Paragraph('UNIT COST', st_hdr_r),
            Paragraph('DISC %', st_hdr_r), Paragraph('VAT %', st_hdr_r),
            Paragraph('LINE TOTAL', st_hdr_r),
        ]]
        subtotal = tax_total = 0.0
        for i, it in enumerate(po.items or [], 1):
            line_total = float(it.get('total') or 0)
            tax = float(it.get('tax_amount') or 0)
            subtotal += line_total
            tax_total += tax
            item_rows.append([
                Paragraph(str(i), st_cell),
                Paragraph(it.get('name') or '', st_cell),
                Paragraph(str(int(it.get('qty') or 0)), st_cell_r),
                Paragraph(money(it.get('unit_cost')), st_cell_r),
                Paragraph(str(float(it.get('discount_percent') or 0)), st_cell_r),
                Paragraph(str(float(it.get('tax_percent') or 0)), st_cell_r),
                Paragraph(money(line_total + tax), st_cell_b_r),
            ])
        if not (po.items or []):
            item_rows.append([Paragraph('No items on this order', st_note), '', '', '', '', '', ''])
        items_tbl = Table(item_rows, colWidths=[
            8 * mm, PAGE_W - (8 + 14 + 26 + 16 + 14 + 28) * mm,
            14 * mm, 26 * mm, 16 * mm, 14 * mm, 28 * mm,
        ], repeatRows=1)
        items_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 0), (-1, 0), 1, TEAL),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 1), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ]))
        story.append(items_tbl)
        story.append(Spacer(1, 5 * mm))

        # ── Totals card ──────────────────────────────────────────────────────
        totals = [
            [Paragraph('Subtotal (excl. VAT)', st_tot_lbl), Paragraph(money(subtotal), st_tot_val)],
            [Paragraph('VAT', st_tot_lbl), Paragraph(money(tax_total), st_tot_val)],
        ]
        if float(po.shipping_cost or 0) > 0:
            totals.append([
                Paragraph('Shipping (excluded from PO total — billed via expenses)', st_tot_lbl),
                Paragraph(money(po.shipping_cost), st_tot_val),
            ])
        totals.append([
            Paragraph('TOTAL (INCL. VAT)', st_tot_lbl_b),
            Paragraph(money(po.total_cost), st_tot_val_b),
        ])
        tot_w = 88 * mm
        totals_tbl = Table(totals, colWidths=[tot_w - 42 * mm, 42 * mm], hAlign='RIGHT')
        totals_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -2), BG),
            ('BOX', (0, 0), (-1, -2), 0.75, BORDER),
            ('BACKGROUND', (0, -1), (-1, -1), TEAL),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('RIGHTPADDING', (0, 0), (-1, -1), 10),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(totals_tbl)

        # ── Notes ───────────────────────────────────────────────────────────
        if po.notes:
            story.append(Spacer(1, 6 * mm))
            story.append(section_heading('NOTES'))
            story.append(Spacer(1, 2.5 * mm))
            story.append(Paragraph(po.notes, st_body))

        # ── Footer (every page): powered by AdhereMed + website link ────────
        def draw_footer(canvas, doc_):
            canvas.saveState()
            canvas.setFillColor(NAVY)
            canvas.rect(0, A4[1] - 4 * mm, A4[0], 4 * mm, stroke=0, fill=1)
            canvas.setFillColor(TEAL)
            canvas.rect(0, A4[1] - 5.2 * mm, A4[0], 1.2 * mm, stroke=0, fill=1)
            canvas.setStrokeColor(BORDER)
            canvas.setLineWidth(0.5)
            canvas.line(15 * mm, 13 * mm, A4[0] - 15 * mm, 13 * mm)
            canvas.setFont('Helvetica', 7.5)
            canvas.setFillColor(SLATE)
            canvas.drawCentredString(A4[0] / 2, 8.5 * mm, 'Powered by AdhereMed')
            canvas.setFont('Helvetica-Bold', 8)
            canvas.setFillColor(TEAL)
            canvas.drawCentredString(A4[0] / 2, 4.5 * mm, 'www.adheremed.co')
            canvas.linkURL(
                'https://www.adheremed.co',
                (A4[0] / 2 - 60, 3.5 * mm, A4[0] / 2 + 60, 7.5 * mm),
            )
            canvas.restoreState()

        doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
        buf.seek(0)
        fname = f'purchase_order_{po.po_number}.pdf'
        return HttpResponse(
            buf.read(),
            content_type='application/pdf',
            headers={'Content-Disposition': f'attachment; filename="{fname}"'},
        )

    @action(detail=False, methods=['get'], url_path='report-pdf')
    def report_pdf(self, request):
        """Generate a branded premium PDF report of purchase orders.

        Respects the same filters as the list endpoint (status, supplier,
        branch) plus optional date_from / date_to on the order date.
        """
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.units import mm
            from reportlab.lib.colors import HexColor
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.lib.enums import TA_RIGHT, TA_CENTER
            from reportlab.platypus import (
                SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
            )
            from django.utils import timezone as _tz
        except ImportError:
            return Response({'detail': 'reportlab not installed.'}, status=500)

        qs = self.filter_queryset(self.get_queryset())
        params = request.query_params
        date_from = params.get('date_from')
        date_to = params.get('date_to')
        pos = list(qs)

        logo_path, business_name = self._resolve_branding(request)

        NAVY = HexColor('#0F172A')
        TEAL = HexColor('#0D9488')
        SLATE = HexColor('#64748B')
        BORDER = HexColor('#E2E8F0')
        BG = HexColor('#F8FAFC')
        WHITE = HexColor('#FFFFFF')
        STATUS_COLORS = {
            'draft': HexColor('#64748B'),
            'sent': HexColor('#2563EB'),
            'received': HexColor('#16A34A'),
            'partial': HexColor('#D97706'),
            'cancelled': HexColor('#DC2626'),
            'returned': HexColor('#7C3AED'),
        }

        def money(v):
            try:
                return f'KSh {float(v or 0):,.2f}'
            except (TypeError, ValueError):
                return 'KSh 0.00'

        def fmt_date(d):
            return d.strftime('%d %b %Y') if d else '—'

        # ── Aggregate KPIs ────────────────────────────────────────────────
        total_spend = 0.0
        vat_total = 0.0
        shipping_total = 0.0
        status_agg = {}
        supplier_agg = {}
        for po in pos:
            spend = float(po.total_cost or 0)
            total_spend += spend
            shipping_total += float(po.shipping_cost or 0)
            for it in (po.items or []):
                vat_total += float(it.get('tax_amount') or 0)
            key = po.status
            if key not in status_agg:
                status_agg[key] = {'count': 0, 'value': 0.0, 'label': po.get_status_display()}
            status_agg[key]['count'] += 1
            status_agg[key]['value'] += spend
            sname = po.supplier.name if po.supplier else '—'
            if sname not in supplier_agg:
                supplier_agg[sname] = 0.0
            supplier_agg[sname] += spend
        open_count = sum(v['count'] for k, v in status_agg.items()
                         if k in ('draft', 'sent', 'partial'))
        received_count = status_agg.get('received', {}).get('count', 0)
        top_suppliers = sorted(supplier_agg.items(), key=lambda x: -x[1])[:5]

        PAGE = landscape(A4)
        PAGE_W = PAGE[0] - 30 * mm

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=PAGE,
            leftMargin=15 * mm, rightMargin=15 * mm,
            topMargin=12 * mm, bottomMargin=18 * mm,
            title='Purchase Orders Report',
        )

        st_brand = ParagraphStyle('brand', fontName='Helvetica-Bold', fontSize=13,
                                   textColor=NAVY, leading=16)
        st_brand_sub = ParagraphStyle('brandsub', fontName='Helvetica', fontSize=8,
                                      textColor=SLATE, leading=11)
        st_doctitle = ParagraphStyle('doctitle', fontName='Helvetica-Bold', fontSize=20,
                                     textColor=TEAL, leading=24, alignment=TA_RIGHT)
        st_docsub = ParagraphStyle('docsub', fontName='Helvetica', fontSize=8.5,
                                   textColor=SLATE, leading=12, alignment=TA_RIGHT)
        st_label = ParagraphStyle('label', fontName='Helvetica-Bold', fontSize=6.5,
                                  textColor=SLATE, leading=9)
        st_value = ParagraphStyle('value', fontName='Helvetica-Bold', fontSize=11,
                                  textColor=NAVY, leading=14)
        st_value_sub = ParagraphStyle('valuesub', fontName='Helvetica', fontSize=7,
                                      textColor=SLATE, leading=9)
        st_section = ParagraphStyle('section', fontName='Helvetica-Bold', fontSize=9,
                                    textColor=NAVY, leading=12)
        st_cell = ParagraphStyle('cell', fontName='Helvetica', fontSize=8, leading=11,
                                 textColor=NAVY)
        st_cell_r = ParagraphStyle('cellr', parent=st_cell, alignment=TA_RIGHT)
        st_cell_b = ParagraphStyle('cellb', parent=st_cell, fontName='Helvetica-Bold')
        st_cell_b_r = ParagraphStyle('cellbr', parent=st_cell_b, alignment=TA_RIGHT)
        st_hdr = ParagraphStyle('hdr', parent=st_cell_b, textColor=WHITE)
        st_hdr_r = ParagraphStyle('hdrr', parent=st_hdr, alignment=TA_RIGHT)
        st_pill = ParagraphStyle('pill', fontName='Helvetica-Bold', fontSize=7.5,
                                 textColor=WHITE, alignment=TA_CENTER)
        st_note = ParagraphStyle('note', fontName='Helvetica-Oblique', fontSize=7.5,
                                 textColor=SLATE, leading=10)

        def section_heading(text, width):
            bar = Table([['']], colWidths=[1.2 * mm], rowHeights=[3.6 * mm])
            bar.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), TEAL)]))
            t = Table([[bar, Paragraph(text, st_section)]], colWidths=[3 * mm, width - 3 * mm])
            t.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
                ('TOPPADDING', (0, 0), (-1, -1), 0),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ]))
            return t

        story = []

        # ── Header ──────────────────────────────────────────────────────────
        logo = RLImage(logo_path, width=15 * mm, height=15 * mm, kind='proportional')
        tenant = getattr(request, 'tenant', None)
        header_subs = []
        if tenant:
            addr = ', '.join(x for x in [getattr(tenant, 'address', '') or '',
                                         getattr(tenant, 'city', '') or ''] if x)
            if addr:
                header_subs.append(addr)
            contact = ' · '.join(x for x in [getattr(tenant, 'phone', '') or '',
                                             getattr(tenant, 'email', '') or ''] if x)
            if contact:
                header_subs.append(contact)
        left_block = [Paragraph(business_name, st_brand)] + \
            [Paragraph(s, st_brand_sub) for s in header_subs]
        date_range = ''
        if date_from or date_to:
            date_range = f'Date range: {date_from or "—"} to {date_to or "—"}'
        right_block = [
            Paragraph('PURCHASE ORDERS REPORT', st_doctitle),
            Paragraph(f'Generated {_tz.now().strftime("%d %b %Y, %H:%M")}', st_docsub),
            *( [Paragraph(date_range, st_docsub)] if date_range else [] ),
        ]
        header_tbl = Table([[logo, left_block, right_block]],
                           colWidths=[17 * mm, PAGE_W - 17 * mm - 80 * mm, 80 * mm])
        header_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, 0), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(header_tbl)
        story.append(Spacer(1, 4 * mm))
        strip = Table([['']], colWidths=[PAGE_W], rowHeights=[1.6 * mm])
        strip.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), TEAL)]))
        story.append(strip)
        story.append(Spacer(1, 5 * mm))

        # ── KPI cards ──────────────────────────────────────────────────────
        def kpi_card(label, value, sub=''):
            rows = [[Paragraph(label, st_label)], [Paragraph(value, st_value)]]
            if sub:
                rows.append([Paragraph(sub, st_value_sub)])
            return rows

        kpis = [
            kpi_card('TOTAL ORDERS', str(len(pos)), f'{open_count} open'),
            kpi_card('RECEIVED', str(received_count), 'stock updated'),
            kpi_card('TOTAL SPEND', money(total_spend), 'incl. VAT'),
            kpi_card('VAT', money(vat_total), 'across all items'),
            kpi_card('SHIPPING', money(shipping_total), 'billed via expenses'),
        ]
        gap = 4 * mm
        kpi_w = (PAGE_W - 4 * gap) / 5
        kpi_tbl = Table(
            [[kpis[0], '', kpis[1], '', kpis[2], '', kpis[3], '', kpis[4]]],
            colWidths=[kpi_w, gap, kpi_w, gap, kpi_w, gap, kpi_w, gap, kpi_w],
        )
        kpi_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BACKGROUND', (0, 0), (0, 0), BG),
            ('BACKGROUND', (2, 0), (2, 0), BG),
            ('BACKGROUND', (4, 0), (4, 0), BG),
            ('BACKGROUND', (6, 0), (6, 0), BG),
            ('BACKGROUND', (8, 0), (8, 0), BG),
            ('BOX', (0, 0), (0, 0), 0.75, BORDER),
            ('BOX', (2, 0), (2, 0), 0.75, BORDER),
            ('BOX', (4, 0), (4, 0), 0.75, BORDER),
            ('BOX', (6, 0), (6, 0), 0.75, BORDER),
            ('BOX', (8, 0), (8, 0), 0.75, BORDER),
            ('LEFTPADDING', (0, 0), (-1, -1), 9),
            ('RIGHTPADDING', (0, 0), (-1, -1), 9),
            ('TOPPADDING', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ]))
        story.append(kpi_tbl)
        story.append(Spacer(1, 6 * mm))

        # ── Status breakdown + Top suppliers (side by side) ────────────────
        half = (PAGE_W - 10 * mm) / 2
        status_rows = [[
            Paragraph('STATUS', st_hdr), Paragraph('ORDERS', st_hdr_r),
            Paragraph('VALUE (KSH)', st_hdr_r),
        ]]
        for key in ('draft', 'sent', 'partial', 'received', 'returned', 'cancelled'):
            agg = status_agg.get(key)
            if not agg:
                continue
            pill = Table([[Paragraph(agg['label'], st_pill)]], colWidths=[22 * mm])
            pill.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), STATUS_COLORS.get(key, SLATE)),
                ('TOPPADDING', (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ]))
            status_rows.append([
                pill,
                Paragraph(str(agg['count']), st_cell_r),
                Paragraph(f"{agg['value']:,.2f}", st_cell_r),
            ])
        status_tbl = Table(status_rows, colWidths=[half - 44 * mm, 22 * mm, 22 * mm])
        status_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        supplier_rows = [[
            Paragraph('SUPPLIER', st_hdr),
            Paragraph('TOTAL SPEND (KSH)', st_hdr_r),
        ]]
        for name, value in top_suppliers:
            supplier_rows.append([
                Paragraph(name, st_cell),
                Paragraph(f'{value:,.2f}', st_cell_r),
            ])
        if not top_suppliers:
            supplier_rows.append([Paragraph('No data', st_note), ''])
        supplier_tbl = Table(supplier_rows, colWidths=[half - 34 * mm, 34 * mm])
        supplier_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        panels = Table([[
            [section_heading('STATUS BREAKDOWN', half), Spacer(1, 3 * mm), status_tbl],
            '',
            [section_heading('TOP SUPPLIERS BY SPEND', half), Spacer(1, 3 * mm), supplier_tbl],
        ]], colWidths=[half, 10 * mm, half])
        panels.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(panels)
        story.append(Spacer(1, 6 * mm))

        # ── PO ledger ──────────────────────────────────────────────────────
        story.append(section_heading(f'PURCHASE ORDERS ({len(pos)})', PAGE_W))
        story.append(Spacer(1, 3.5 * mm))
        ledger_rows = [[
            Paragraph('#', st_hdr), Paragraph('PO NUMBER', st_hdr),
            Paragraph('SUPPLIER', st_hdr), Paragraph('STATUS', st_hdr),
            Paragraph('ORDER DATE', st_hdr_r), Paragraph('EXPECTED', st_hdr_r),
            Paragraph('ITEMS', st_hdr_r), Paragraph('SHIPPING (KSH)', st_hdr_r),
            Paragraph('TOTAL (KSH)', st_hdr_r),
        ]]
        for i, po in enumerate(pos, 1):
            pill = Table([[Paragraph(po.get_status_display(), st_pill)]], colWidths=[26 * mm])
            pill.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), STATUS_COLORS.get(po.status, SLATE)),
                ('TOPPADDING', (0, 0), (-1, -1), 2),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ]))
            ledger_rows.append([
                Paragraph(str(i), st_cell),
                Paragraph(po.po_number, st_cell_b),
                Paragraph(po.supplier.name if po.supplier else '—', st_cell),
                pill,
                Paragraph(fmt_date(po.order_date), st_cell_r),
                Paragraph(fmt_date(po.expected_delivery), st_cell_r),
                Paragraph(str(len(po.items or [])), st_cell_r),
                Paragraph(f"{float(po.shipping_cost or 0):,.2f}", st_cell_r),
                Paragraph(f"{float(po.total_cost or 0):,.2f}", st_cell_b_r),
            ])
        if not pos:
            ledger_rows.append([Paragraph('No purchase orders found for the selected filters.', st_note), '', '', '', '', '', '', '', ''])
        ledger_w = [8 * mm, 34 * mm, 58 * mm, 30 * mm, 26 * mm, 26 * mm, 16 * mm, 30 * mm, 32 * mm]
        ledger_tbl = Table(ledger_rows, colWidths=ledger_w, repeatRows=1)
        ledger_tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), NAVY),
            ('LINEBELOW', (0, 0), (-1, 0), 1, TEAL),
            ('LINEBELOW', (0, 1), (-1, -1), 0.4, BORDER),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [None, BG]),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 4.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4.5),
        ]))
        story.append(ledger_tbl)

        # ── Footer (every page) ─────────────────────────────────────────────
        def draw_footer(canvas, doc_):
            canvas.saveState()
            canvas.setFillColor(NAVY)
            canvas.rect(0, PAGE[1] - 4 * mm, PAGE[0], 4 * mm, stroke=0, fill=1)
            canvas.setFillColor(TEAL)
            canvas.rect(0, PAGE[1] - 5.2 * mm, PAGE[0], 1.2 * mm, stroke=0, fill=1)
            canvas.setStrokeColor(BORDER)
            canvas.setLineWidth(0.5)
            canvas.line(15 * mm, 13 * mm, PAGE[0] - 15 * mm, 13 * mm)
            canvas.setFont('Helvetica', 7.5)
            canvas.setFillColor(SLATE)
            canvas.drawCentredString(PAGE[0] / 2, 8.5 * mm, 'Powered by AdhereMed')
            canvas.setFont('Helvetica-Bold', 8)
            canvas.setFillColor(TEAL)
            canvas.drawCentredString(PAGE[0] / 2, 4.5 * mm, 'www.adheremed.co')
            canvas.linkURL(
                'https://www.adheremed.co',
                (PAGE[0] / 2 - 60, 3.5 * mm, PAGE[0] / 2 + 60, 7.5 * mm),
            )
            canvas.restoreState()

        doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
        buf.seek(0)
        stamp = _tz.now().strftime('%Y%m%d')
        fname = f'purchase_orders_report_{stamp}.pdf'
        return HttpResponse(
            buf.read(),
            content_type='application/pdf',
            headers={'Content-Disposition': f'attachment; filename="{fname}"'},
        )

    @action(detail=False, methods=['get'], url_path='export')
    def export(self, request):
        """Export purchase orders as .xlsx or .csv (?fmt=excel|csv).

        Accepts the same filters as the list endpoint (status, supplier, branch).
        """
        qs = self.filter_queryset(self.get_queryset())
        fmt = (request.query_params.get('fmt') or 'excel').lower()
        headers_row = ['PO Number', 'Supplier', 'Branch', 'Status',
                       'Order Date', 'Expected Delivery', 'Items',
                       'Total Cost (KSh)', 'Notes']

        def _po_row(po):
            return [
                po.po_number,
                po.supplier.name if po.supplier else '',
                po.branch.name if po.branch else '',
                po.get_status_display(),
                po.order_date.strftime('%Y-%m-%d') if po.order_date else '',
                po.expected_delivery.strftime('%Y-%m-%d') if po.expected_delivery else '',
                len(po.items or []),
                float(po.total_cost or 0),
                po.notes or '',
            ]

        from django.utils import timezone as _tz
        stamp = _tz.now().strftime('%Y%m%d')

        if fmt == 'csv':
            import csv as _csv
            fname = f'purchase_orders_{stamp}.csv'
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="{fname}"'
            writer = _csv.writer(response)
            writer.writerow(headers_row)
            for po in qs:
                writer.writerow(_po_row(po))
            return response

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response({'detail': 'openpyxl not installed.'}, status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Purchase Orders'
        ws.append(headers_row)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='0D9488')
        for po in qs:
            ws.append(_po_row(po))
        for i, w in enumerate([18, 28, 20, 18, 14, 18, 8, 18, 40], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = f'purchase_orders_{stamp}.xlsx'
        return HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{fname}"'},
        )

    @action(detail=False, methods=['post'], url_path='export-items')
    def export_items(self, request):
        """Export a list of PO items as an .xlsx in the import-template format.

        Payload: { "items": [{name, qty, unit_cost, unit_selling_price,
                              discount_percent, tax_percent, batch_number,
                              expiry_date}, ...] }
        Used to export the current (possibly unsaved) items of a PO form.
        """
        items = request.data.get('items')
        if not isinstance(items, list) or not items:
            return Response({'detail': 'items must be a non-empty list.'}, status=400)

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response({'detail': 'openpyxl not installed.'}, status=500)

        headers = ['Item Name', 'Quantity', 'Unit Cost (KSh)',
                   'Unit Selling Price (KSh)', 'Discount %', 'VAT %',
                   'Batch #', 'Expiry Date']
        wb = Workbook()
        ws = wb.active
        ws.title = 'PO Items'
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='0D9488')

        for it in items:
            ws.append([
                str(it.get('name') or it.get('medication_name') or ''),
                it.get('qty') or 0,
                float(it.get('unit_cost') or 0),
                float(it.get('unit_selling_price') or 0),
                float(it.get('discount_percent') or 0),
                float(it.get('tax_percent') or 0),
                str(it.get('batch_number') or ''),
                str(it.get('expiry_date') or ''),
            ])
        for i, w in enumerate([32, 12, 16, 22, 12, 10, 14, 14], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename="po_items.xlsx"'},
        )

    @action(detail=False, methods=['post'], url_path='import-items-preview')
    def import_items_preview(self, request):
        """Parse an uploaded spreadsheet of PO items for review.

        Nothing is saved — rows are added to the client-side form after review.
        """
        uploaded = request.FILES.get('file')
        if not uploaded:
            return Response({'detail': 'No file provided.'}, status=400)
        try:
            parsed = self._parse_item_spreadsheet(uploaded)
        except ValueError as exc:
            return Response({'detail': str(exc)}, status=400)

        from django.db.models import Sum, Q
        from django.db.models.functions import Coalesce
        from inventory.models import MedicationStock

        branch_id = request.query_params.get('branch')
        stocks_by_name = {}
        stocks = (MedicationStock.objects
                  .annotate(qty=Coalesce(Sum('batches__quantity_remaining',
                                             filter=Q(batches__quantity_remaining__gt=0)), 0))
                  .only('id', 'medication_name', 'cost_price', 'selling_price', 'branch'))
        for s in stocks:
            stocks_by_name.setdefault((s.medication_name or '').lower(), []).append(s)

        def find_stock(name):
            cands = stocks_by_name.get(name.lower())
            if not cands:
                return None
            if branch_id:
                for c in cands:
                    if str(c.branch_id) == str(branch_id):
                        return c
            return cands[0]

        rows, counts = [], {'matched': 0, 'new': 0, 'error': 0}
        for num, raw in parsed:
            errors = {}
            name = str(raw.get('name') or '').strip()
            if not name:
                errors['name'] = 'Required'

            qty, qty_err = self._clean_number(raw.get('qty'))
            if raw.get('qty') is None:
                errors['qty'] = 'Required'
            elif qty_err:
                errors['qty'] = qty_err
            elif qty <= 0:
                errors['qty'] = 'Must be > 0'
            else:
                qty = int(qty)

            stock = find_stock(name) if name else None

            unit_cost, cost_err = self._clean_number(
                raw.get('unit_cost'),
                float(stock.cost_price or 0) if stock else 0,
            )
            if cost_err:
                errors['unit_cost'] = cost_err
            sell_default = float(stock.selling_price or 0) if stock else 0
            unit_selling_price, sell_err = self._clean_number(raw.get('unit_selling_price'), sell_default)
            if sell_err:
                errors['unit_selling_price'] = sell_err

            discount_percent, disc_err = self._clean_number(raw.get('discount_percent'))
            if disc_err:
                errors['discount_percent'] = disc_err
            tax_percent, tax_err = self._clean_number(raw.get('tax_percent'))
            if tax_err:
                errors['tax_percent'] = tax_err

            row_status = 'error' if errors else ('matched' if stock else 'new')
            counts[row_status] += 1
            rows.append({
                'row_num': num,
                'name': name,
                'stock_id': stock.id if stock else None,
                'current_stock': int(stock.qty) if stock else None,
                'qty': int(qty) if not errors.get('qty') else qty,
                'unit_cost': unit_cost,
                'unit_selling_price': unit_selling_price,
                'discount_percent': discount_percent,
                'tax_percent': tax_percent,
                'batch_number': str(raw.get('batch_number') or '').strip(),
                'expiry_date': self._clean_date(raw.get('expiry_date')),
                'status': row_status,
                'errors': errors,
            })

        return Response({
            'rows': rows,
            'summary': {
                'total': len(rows),
                'matched': counts['matched'],
                'new': counts['new'],
                'error': counts['error'],
            },
        })

    @action(detail=False, methods=['get'], url_path='import-items-template')
    def import_items_template(self, request):
        """Download an .xlsx template for importing PO items."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response({'detail': 'openpyxl not installed.'}, status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = 'PO Items'
        headers = ['Item Name', 'Quantity', 'Unit Cost (KSh)',
                   'Unit Selling Price (KSh)', 'Discount %', 'VAT %',
                   'Batch #', 'Expiry Date']
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='0D9488')
        ws.append(['Paracetamol 500mg', 100, 12.5, 20, 0, 16, 'B-2201', '2027-06-30'])
        ws.append(['Amoxicillin 250mg', 50, 45, 75, 5, 0, '', ''])
        for i, w in enumerate([32, 12, 16, 22, 12, 10, 14, 14], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fname = 'po_items_import_template.xlsx'
        return HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{fname}"'},
        )


class GoodsReceivedNoteViewSet(viewsets.ModelViewSet):
    queryset = GoodsReceivedNote.objects.select_related('purchase_order', 'received_by').all()
    serializer_class = GoodsReceivedNoteSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['purchase_order']
    search_fields = ['grn_number']
    ordering_fields = ['received_date']
